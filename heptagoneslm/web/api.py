"""Logique de l'interface web : rapports à la demande et exports.

Indépendant du serveur HTTP (testable directement). Chaque rapport renvoie un
dictionnaire sérialisable en JSON ::

    {
      "type": "...", "title": "...", "description": "...",
      "columns": [{"key": ..., "label": ..., "type": "text|int|float|pct|hours|date|datetime"}],
      "rows": [{...}],
      "chart": None | {"kind": "line"|"bar"|"heatmap", ...},
      "kpis": [{"label": ..., "value": ..., "type": ...}],
      "notes": ["..."],
    }

Les mêmes résultats alimentent l'affichage et les exports (Excel, CSV, JSON) :
on exporte exactement ce qui est affiché, avec les mêmes filtres.
"""

from __future__ import annotations

import csv
import io
import json
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Any, Callable

from heptagoneslm.reporting.dataset import WEEKDAYS_FR, _server_filter, build_report_data, load_features
from heptagoneslm.reporting.sanitize import csv_text, excel_row
from heptagoneslm.storage.db import Database

GRANULARITIES = ("hour", "day", "week", "month")
METRICS = ("peak", "avg", "utilization")
DIMENSIONS = {
    "user": "Utilisateur",
    "host": "Poste",
    "department": "Service",
    "product": "Produit",
    "server": "Serveur",
    "publisher": "Éditeur",
    "site": "Site",
}
MAX_CHART_SERIES = 8   # au-delà, les séries restent dans le tableau (règle de lisibilité)
MAX_BAR_ITEMS = 15
NOT_SET = "(non renseigné)"


class ParamError(ValueError):
    """Paramètre de requête invalide."""


@dataclass
class ReportParams:
    start: datetime
    end: datetime                      # exclu
    servers: list[str] = field(default_factory=list)
    products: list[str] = field(default_factory=list)
    users: list[str] = field(default_factory=list)
    departments: list[str] = field(default_factory=list)
    sites: list[str] = field(default_factory=list)
    granularity: str = "day"
    metric: str = "peak"
    dimension: str = "user"
    business_hours: tuple[int, int] = (8, 19)
    costs: dict[str, float] = field(default_factory=dict)   # coût annuel d'un jeton par produit
    currency: str = "€"
    target_coverage: float = 0.95
    scoped: bool = False               # périmètre imposé par le rôle (responsable de service)

    def describe(self) -> str:
        parts = [f"Période du {self.start:%d/%m/%Y} au {(self.end - timedelta(seconds=1)):%d/%m/%Y}"]
        for label, values in (("serveurs", self.servers), ("produits", self.products),
                              ("utilisateurs", self.users), ("services", self.departments), ("sites", self.sites)):
            if values:
                parts.append(f"{label} : {', '.join(values)}")
        return " — ".join(parts)


def params_from_config(query: dict[str, list[str]], config, today: date | None = None) -> "ReportParams":
    """Paramètres de requête + réglages de la configuration (heures ouvrées, coûts…)."""
    p = parse_params(query, today=today, business_hours=config.reporting.business_hours)
    p.costs = config.product_costs()
    p.currency = config.reporting.currency
    p.target_coverage = config.reporting.target_coverage
    return p


def parse_params(query: dict[str, list[str]], today: date | None = None,
                 business_hours: tuple[int, int] = (8, 19)) -> ReportParams:
    """Construit les paramètres à partir d'une query string (clés répétables).

    ``start`` et ``end`` sont des dates AAAA-MM-JJ, ``end`` **incluse**.
    Par défaut : les 30 derniers jours, aujourd'hui inclus.
    """
    today = today or date.today()

    def one(key: str, default: str | None = None) -> str | None:
        values = [v for v in query.get(key, []) if v != ""]
        return values[0] if values else default

    def many(key: str) -> list[str]:
        out: list[str] = []
        for v in query.get(key, []):
            out.extend(x for x in v.split("|") if x)
        return out

    try:
        end_day = date.fromisoformat(one("end") or today.isoformat())
        start_day = date.fromisoformat(one("start") or (end_day - timedelta(days=29)).isoformat())
        end_day + timedelta(days=1)  # la fin exclue doit rester représentable
    except (ValueError, OverflowError):
        raise ParamError("Dates attendues au format AAAA-MM-JJ (entre 0001-01-02 et 9999-12-30)") from None
    if start_day > end_day:
        raise ParamError("La date de début doit précéder la date de fin")
    if (end_day - start_day).days > 3660:
        raise ParamError("Période limitée à 10 ans")

    granularity = one("granularity", "day")
    metric = one("metric", "peak")
    dimension = one("dimension", "user")
    if granularity not in GRANULARITIES:
        raise ParamError(f"granularity doit valoir {', '.join(GRANULARITIES)}")
    if metric not in METRICS:
        raise ParamError(f"metric doit valoir {', '.join(METRICS)}")
    if dimension not in DIMENSIONS:
        raise ParamError(f"dimension doit valoir {', '.join(DIMENSIONS)}")

    return ReportParams(
        start=datetime.combine(start_day, datetime.min.time()),
        end=datetime.combine(end_day + timedelta(days=1), datetime.min.time()),
        servers=many("server"),
        products=many("product"),
        users=[u.lower() for u in many("user")],
        departments=many("department"),
        sites=many("site"),
        granularity=granularity,
        metric=metric,
        dimension=dimension,
        business_hours=business_hours,
    )


# ---------------------------------------------------------------------------
# Métadonnées (listes des filtres, état des collectes)
# ---------------------------------------------------------------------------


def meta(db: Database) -> dict:
    servers = [r[0] for r in db.query("SELECT name FROM license_server ORDER BY name")]
    products = db.query_dicts(
        "SELECT DISTINCT COALESCE(f.product, f.name) AS product, f.publisher, srv.name AS server "
        "FROM feature f JOIN license_server srv ON srv.id = f.server_id ORDER BY 1"
    )
    by_product: dict[str, dict] = {}
    for p in products:
        entry = by_product.setdefault(p["product"], {"name": p["product"], "publisher": p["publisher"], "servers": []})
        if p["server"] not in entry["servers"]:
            entry["servers"].append(p["server"])
    users = db.query_dicts("SELECT username, display_name, department FROM app_user ORDER BY username")
    departments = sorted({u["department"] for u in users if u["department"]})
    sites = sorted({r[0] for r in db.query("SELECT DISTINCT site FROM workstation WHERE site IS NOT NULL")})
    first = db.scalar("SELECT MIN(hour_start) FROM usage_hourly")
    first_session = db.scalar("SELECT MIN(checkout_time) FROM license_session")
    last = db.scalar("SELECT MAX(hour_start) FROM usage_hourly")
    last_session = db.scalar("SELECT MAX(COALESCE(checkin_time, last_seen)) FROM license_session")
    starts = [d for d in (first, first_session) if d]
    ends = [d for d in (last, last_session) if d]
    return {
        "servers": servers,
        "products": list(by_product.values()),
        "users": [
            {"username": u["username"], "display_name": u["display_name"], "department": u["department"]}
            for u in users
        ],
        "departments": departments,
        "sites": sites,
        "data_start": _iso(min(starts)) if starts else None,
        "data_end": _iso(max(ends)) if ends else None,
        "collections": collections_status(db),
        "dimensions": DIMENSIONS,
    }


def collections_status(db: Database) -> list[dict]:
    """Dernière exécution de chaque collecteur, avec le nombre d'échecs sur 24 h."""
    since = datetime.now() - timedelta(hours=24)
    rows = db.query_dicts(
        "SELECT srv.name AS server, r.collector, r.started_at, r.status, r.items, r.message "
        "FROM collection_run r JOIN license_server srv ON srv.id = r.server_id "
        "WHERE r.id IN (SELECT MAX(id) FROM collection_run GROUP BY server_id, collector) "
        "ORDER BY srv.name, r.collector"
    )
    errors = dict(
        ((r[0], r[1]), r[2])
        for r in db.query(
            "SELECT srv.name, r.collector, COUNT(*) FROM collection_run r "
            "JOIN license_server srv ON srv.id = r.server_id "
            "WHERE r.status = 'ERROR' AND r.started_at >= ? GROUP BY srv.name, r.collector",
            [since],
        )
    )
    return [
        {
            "server": r["server"],
            "collector": r["collector"],
            "last_run": _iso(r["started_at"]),
            "status": r["status"],
            "items": r["items"],
            "message": r["message"],
            "errors_24h": errors.get((r["server"], r["collector"]), 0),
        }
        for r in rows
    ]


# ---------------------------------------------------------------------------
# Rapports
# ---------------------------------------------------------------------------


def _iso(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    return value


def _col(key: str, label: str, type_: str = "text") -> dict:
    return {"key": key, "label": label, "type": type_}


def _result(type_: str, title: str, description: str, columns: list[dict], rows: list[dict],
            chart: dict | None = None, kpis: list[dict] | None = None, notes: list[str] | None = None) -> dict:
    return {
        "type": type_,
        "title": title,
        "description": description,
        "columns": columns,
        "rows": [{k: _iso(v) for k, v in r.items()} for r in rows],
        "chart": chart,
        "kpis": kpis or [],
        "notes": notes or [],
    }


def _data(db: Database, p: ReportParams):
    return build_report_data(db, p.start, p.end, servers=p.servers or None,
                             business_hours=p.business_hours, products=p.products or None)


def _filter_sessions(sessions: list[dict], p: ReportParams) -> list[dict]:
    out = sessions
    if p.users:
        wanted = set(p.users)
        out = [s for s in out if s["username"] in wanted]
    if p.departments:
        wanted = set(p.departments)
        out = [s for s in out if (s["department"] or NOT_SET) in wanted]
    if p.sites:
        wanted = set(p.sites)
        out = [s for s in out if (s.get("site") or NOT_SET) in wanted]
    return out


def visible_usernames(db: Database, p: ReportParams) -> set[str] | None:
    """Utilisateurs visibles avec les filtres utilisateur / service (None : tous)."""
    if not p.users and not p.departments:
        return None
    rows = db.query_dicts("SELECT username, department FROM app_user")
    wanted_deps = set(p.departments)
    wanted_users = set(p.users)
    return {r["username"] for r in rows
            if (not wanted_users or r["username"] in wanted_users)
            and (not wanted_deps or (r["department"] or NOT_SET) in wanted_deps)}


def _user_filter_note(p: ReportParams) -> list[str]:
    if p.users or p.departments or p.sites:
        return ["Les filtres utilisateur/service/site ne s'appliquent pas aux taux d'occupation, "
                "qui sont des compteurs globaux du serveur de licences."]
    return []


def report_summary(db: Database, p: ReportParams) -> dict:
    data = _data(db, p)
    sessions = _filter_sessions(data.sessions, p)
    rows = []
    for s in sorted(data.summaries, key=lambda s: (-(s.avg_utilization or 0), s.label)):
        rows.append({
            "server": s.server, "publisher": s.publisher, "product": s.label, "feature": s.feature,
            "issued": s.issued, "peak": s.peak_in_use, "avg": round(s.avg_in_use, 2),
            "avg_utilization": s.avg_utilization, "peak_utilization": s.peak_utilization,
            "business_utilization": s.business_utilization, "saturated_hours": s.saturated_hours,
            "denials": s.denials, "users": s.distinct_users, "sessions": s.sessions,
            "hours": round(s.usage_hours, 1),
        })
    issued = sum(s.issued or 0 for s in data.summaries)
    avg_sum = sum(s.avg_in_use for s in data.summaries if s.issued)
    kpis = [
        {"label": "Produits suivis", "value": len(data.summaries), "type": "int"},
        {"label": "Taux d'occupation moyen", "value": (avg_sum / issued) if issued else None, "type": "pct"},
        {"label": "Heures saturées", "value": sum(s.saturated_hours for s in data.summaries), "type": "int"},
        {"label": "Refus", "value": sum(s.denials for s in data.summaries), "type": "int"},
        {"label": "Utilisateurs actifs", "value": len({s["username"] for s in sessions}), "type": "int"},
        {"label": "Heures d'utilisation", "value": round(sum(s["hours_in_period"] for s in sessions), 1), "type": "hours"},
    ]
    top = [r for r in rows if r["avg_utilization"] is not None][:MAX_BAR_ITEMS]
    chart = {
        "kind": "bar", "format": "pct", "title": "Taux d'occupation moyen par produit",
        "items": [{"label": r["product"], "value": r["avg_utilization"]} for r in top],
    } if top else None
    columns = [
        _col("server", "Serveur"), _col("publisher", "Éditeur"), _col("product", "Produit"),
        _col("feature", "Feature"), _col("issued", "Jetons émis", "int"), _col("peak", "Pic", "int"),
        _col("avg", "Moyenne", "float"), _col("avg_utilization", "Taux moyen", "pct"),
        _col("peak_utilization", "Taux au pic", "pct"),
        _col("business_utilization", "Taux heures ouvrées", "pct"),
        _col("saturated_hours", "Heures saturées", "int"), _col("denials", "Refus", "int"),
        _col("users", "Utilisateurs", "int"), _col("sessions", "Sessions", "int"),
        _col("hours", "Heures d'utilisation", "hours"),
    ]
    return _result("summary", "Synthèse par produit", p.describe(), columns, rows, chart, kpis,
                   _user_filter_note(p))


def _bucket(dt: datetime, granularity: str) -> str:
    if granularity == "hour":
        return dt.strftime("%Y-%m-%d %H:00")
    if granularity == "day":
        return dt.strftime("%Y-%m-%d")
    if granularity == "week":
        return (dt.date() - timedelta(days=dt.weekday())).isoformat()
    return dt.strftime("%Y-%m")


def report_timeseries(db: Database, p: ReportParams) -> dict:
    features = {f["id"]: f for f in load_features(db, p.start, p.end, p.servers or None, p.products or None)}
    sql_srv, p_srv = _server_filter(p.servers or None)
    rows_db = db.query(
        "SELECT h.feature_id, h.hour_start, h.samples, h.sum_in_use, h.max_in_use, h.issued "
        "FROM usage_hourly h JOIN feature f ON f.id = h.feature_id "
        "JOIN license_server srv ON srv.id = f.server_id "
        "WHERE h.hour_start >= ? AND h.hour_start < ?" + sql_srv,
        [p.start, p.end, *p_srv],
    )
    # (bucket, feature) -> [samples, sum, max, issued]
    agg: dict[tuple[str, int], list] = {}
    for fid, hour_start, samples, sum_in_use, max_in_use, issued in rows_db:
        if fid not in features or not samples:
            continue
        a = agg.setdefault((_bucket(hour_start, p.granularity), fid), [0, 0, 0, None])
        a[0] += samples
        a[1] += sum_in_use
        a[2] = max(a[2], max_in_use)
        if issued is not None:
            a[3] = max(a[3] or 0, issued)

    rows = []
    for (bucket, fid), (samples, total, peak, issued) in sorted(agg.items(), key=lambda kv: (kv[0][0], features[kv[0][1]]["label"])):
        avg = total / samples
        rows.append({
            "period": bucket, "server": features[fid]["server"], "product": features[fid]["label"],
            "issued": issued, "peak": peak, "avg": round(avg, 2),
            "utilization": (avg / issued) if issued else None,
            "peak_utilization": (peak / issued) if issued else None,
        })

    metric_key = {"peak": "peak", "avg": "avg", "utilization": "utilization"}[p.metric]
    labels = sorted({r["product"] for r in rows})
    # Séries retenues : les plus utilisées en moyenne (lisibilité), ordre stable.
    weight: dict[str, float] = defaultdict(float)
    for r in rows:
        weight[r["product"]] += r["avg"]
    shown = sorted(labels, key=lambda lab: -weight[lab])[:MAX_CHART_SERIES]
    shown.sort()
    x = sorted({r["period"] for r in rows})
    index = {(r["period"], r["product"]): r[metric_key] for r in rows}
    chart = {
        "kind": "line",
        "title": {"peak": "Pic de jetons utilisés", "avg": "Jetons utilisés en moyenne",
                  "utilization": "Taux d'occupation moyen"}[p.metric],
        "format": "pct" if p.metric == "utilization" else ("int" if p.metric == "peak" else "float"),
        "x": x,
        "xType": p.granularity,
        "series": [{"name": lab, "values": [index.get((b, lab)) for b in x]} for lab in shown],
    } if rows else None
    notes = _user_filter_note(p)
    if len(labels) > len(shown):
        notes.append(f"Graphique limité aux {len(shown)} produits les plus utilisés "
                     f"({len(labels)} au total) : filtrez par produit pour les autres. Le tableau contient tout.")
    period_label = {"hour": "Heure", "day": "Jour", "week": "Semaine du", "month": "Mois"}[p.granularity]
    columns = [
        _col("period", period_label), _col("server", "Serveur"), _col("product", "Produit"),
        _col("issued", "Jetons émis", "int"), _col("peak", "Pic", "int"), _col("avg", "Moyenne", "float"),
        _col("utilization", "Taux moyen", "pct"), _col("peak_utilization", "Taux au pic", "pct"),
    ]
    return _result("timeseries", "Usage dans le temps", p.describe(), columns, rows, chart, notes=notes)


def _dimension_value(s: dict, dimension: str) -> str:
    key = {"user": "username", "host": "hostname", "department": "department",
           "product": "product", "server": "server", "publisher": "publisher", "site": "site"}[dimension]
    return s.get(key) or NOT_SET


def report_breakdown(db: Database, p: ReportParams) -> dict:
    data = _data(db, p)
    sessions = _filter_sessions(data.sessions, p)
    groups: dict[str, dict] = {}
    for s in sessions:
        g = groups.setdefault(_dimension_value(s, p.dimension),
                              {"hours": 0.0, "sessions": 0, "users": set(), "products": set(), "hosts": set()})
        g["hours"] += s["hours_in_period"]
        g["sessions"] += 1
        g["users"].add(s["username"])
        g["products"].add(s["product"])
        g["hosts"].add(s["hostname"])
    # Refacturation : coût de chaque produit sur la période, réparti au prorata des heures.
    if p.costs:
        years = max((p.end - p.start).days, 1) / 365.0
        issued_by_product: dict[str, int] = defaultdict(int)
        for s in data.summaries:
            issued_by_product[s.product] += s.issued or 0
        hours_by_product: dict[str, float] = defaultdict(float)
        for s in sessions:
            hours_by_product[s["product"]] += s["hours_in_period"]
        for s in sessions:
            unit = p.costs.get(s["product"])
            if unit is None or not hours_by_product[s["product"]]:
                continue
            share = s["hours_in_period"] / hours_by_product[s["product"]]
            g = groups[_dimension_value(s, p.dimension)]
            g["cost"] = g.get("cost", 0.0) + unit * issued_by_product[s["product"]] * years * share
    total = sum(g["hours"] for g in groups.values()) or 1.0
    rows = [
        {"key": k, "hours": round(g["hours"], 2), "share": g["hours"] / total, "sessions": g["sessions"],
         "users": len(g["users"]), "products": len(g["products"]), "hosts": len(g["hosts"]),
         "avg_session": round(g["hours"] / g["sessions"], 2) if g["sessions"] else None,
         "cost": round(g["cost"], 2) if "cost" in g else None}
        for k, g in sorted(groups.items(), key=lambda kv: -kv[1]["hours"])
    ]
    label = DIMENSIONS[p.dimension]
    chart = {
        "kind": "bar", "format": "hours", "title": f"Heures d'utilisation par {label.lower()}",
        "items": [{"label": r["key"], "value": r["hours"]} for r in rows[:MAX_BAR_ITEMS]],
    } if rows else None
    notes = []
    if len(rows) > MAX_BAR_ITEMS:
        notes.append(f"Graphique limité aux {MAX_BAR_ITEMS} premiers ({len(rows)} au total).")
    columns = [
        _col("key", label), _col("hours", "Heures", "hours"), _col("share", "Part", "pct"),
        _col("sessions", "Sessions", "int"), _col("users", "Utilisateurs", "int"),
        _col("products", "Produits", "int"), _col("hosts", "Postes", "int"),
        _col("avg_session", "Durée moy. session (h)", "hours"),
    ]
    if p.costs:
        columns.append(_col("cost", "Coût refacturé (période)", "money"))
    if p.costs:
        notes.append("Coût refacturé : coût annuel des jetons de chaque produit ramené à la période, "
                     "réparti au prorata des heures d'utilisation.")
    result = _result("breakdown", f"Répartition par {label.lower()}", p.describe(), columns, rows, chart, notes=notes)
    result["currency"] = p.currency
    return result


def report_matrix(db: Database, p: ReportParams) -> dict:
    data = _data(db, p)
    sessions = _filter_sessions(data.sessions, p)
    dimension = "user" if p.dimension == "product" else p.dimension
    products = sorted({s["product"] for s in sessions})
    cells: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for s in sessions:
        cells[_dimension_value(s, dimension)][s["product"]] += s["hours_in_period"]
    rows = []
    for key, per in sorted(cells.items(), key=lambda kv: -sum(kv[1].values())):
        row: dict[str, Any] = {"key": key}
        for i, prod in enumerate(products):
            row[f"p{i}"] = round(per[prod], 2) if prod in per else None
        row["total"] = round(sum(per.values()), 2)
        rows.append(row)
    label = DIMENSIONS[dimension]
    columns = ([_col("key", label)] + [_col(f"p{i}", prod, "hours") for i, prod in enumerate(products)]
               + [_col("total", "Total (h)", "hours")])
    return _result("matrix", f"Croisement {label.lower()} × produit (heures)", p.describe(), columns, rows,
                   {"kind": "matrix"} if rows else None)


def report_heatmap(db: Database, p: ReportParams) -> dict:
    data = _data(db, p)
    rows = []
    values = []
    for hour in range(24):
        row: dict[str, Any] = {"hour": f"{hour:02d}h"}
        line = []
        for wd in range(7):
            v = data.hour_profile.get((wd, hour))
            row[f"d{wd}"] = v
            line.append(v)
        rows.append(row)
        values.append(line)
    has_data = any(v is not None for line in values for v in line)
    columns = [_col("hour", "Heure")] + [_col(f"d{i}", day, "pct") for i, day in enumerate(WEEKDAYS_FR)]
    chart = {"kind": "heatmap", "format": "pct", "title": "Taux d'occupation moyen par heure et jour",
             "rows": [r["hour"] for r in rows], "cols": WEEKDAYS_FR, "values": values} if has_data else None
    return _result("heatmap", "Profil horaire", p.describe(), columns, rows if has_data else [], chart,
                   notes=_user_filter_note(p))


def report_contention(db: Database, p: ReportParams) -> dict:
    data = _data(db, p)
    rows = [dict(c) for c in data.contention]
    columns = [
        _col("date", "Date", "date"), _col("server", "Serveur"), _col("product", "Produit"),
        _col("feature", "Feature"), _col("issued", "Jetons émis", "int"), _col("peak_in_use", "Pic", "int"),
        _col("saturated_hours", "Heures saturées", "int"), _col("denials", "Refus", "int"),
    ]
    by_product: dict[str, int] = defaultdict(int)
    for r in rows:
        by_product[r["product"]] += r["saturated_hours"]
    items = [{"label": k, "value": v} for k, v in sorted(by_product.items(), key=lambda kv: -kv[1]) if v][:MAX_BAR_ITEMS]
    chart = {"kind": "bar", "format": "int", "title": "Heures saturées par produit", "items": items} if items else None
    return _result("contention", "Contention (saturation et refus)", p.describe(), columns, rows, chart)


def report_denials(db: Database, p: ReportParams) -> dict:
    data = _data(db, p)
    rows = data.denials
    visible = visible_usernames(db, p)
    if visible is not None:
        rows = [d for d in rows if (d["username"] or "") in visible]
    columns = [
        _col("event_time", "Horodatage", "datetime"), _col("server", "Serveur"), _col("product", "Produit"),
        _col("feature", "Feature"), _col("username", "Utilisateur"), _col("hostname", "Poste"),
        _col("detail", "Motif"),
    ]
    by_product: dict[str, int] = defaultdict(int)
    for d in rows:
        by_product[d["product"]] += 1
    items = [{"label": k, "value": v} for k, v in sorted(by_product.items(), key=lambda kv: -kv[1])][:MAX_BAR_ITEMS]
    chart = {"kind": "bar", "format": "int", "title": "Refus par produit", "items": items} if items else None
    notes = [] if rows else ["Les refus ne sont connus que pour les serveurs dont le log est collecté (collecteur « log »)."]
    return _result("denials", "Refus de licence", p.describe(), columns, rows, chart, notes=notes)


def report_sessions(db: Database, p: ReportParams) -> dict:
    data = _data(db, p)
    rows = [dict(r, borrowed="Oui" if r.get("borrowed") else "") for r in _filter_sessions(data.sessions, p)]
    columns = [
        _col("server", "Serveur"), _col("publisher", "Éditeur"), _col("product", "Produit"),
        _col("feature", "Feature"), _col("username", "Utilisateur"), _col("display_name", "Nom"),
        _col("department", "Service"), _col("hostname", "Poste"), _col("site", "Site"),
        _col("borrowed", "Emprunt"), _col("checkout_time", "Début", "datetime"),
        _col("checkin_time", "Fin", "datetime"), _col("licenses", "Jetons", "int"),
        _col("hours_in_period", "Heures (période)", "hours"), _col("source", "Source"),
        _col("close_reason", "Motif de fin"),
    ]
    return _result("sessions", "Sessions", p.describe(), columns, rows)


def report_collections(db: Database, p: ReportParams) -> dict:
    rows = db.query_dicts(
        "SELECT r.started_at, srv.name AS server, r.collector, r.status, r.items, r.message "
        "FROM collection_run r JOIN license_server srv ON srv.id = r.server_id "
        "WHERE r.started_at >= ? AND r.started_at < ? ORDER BY r.started_at DESC",
        [p.start, p.end],
    )
    if p.servers:
        rows = [r for r in rows if r["server"] in set(p.servers)]
    columns = [
        _col("started_at", "Début", "datetime"), _col("server", "Serveur"), _col("collector", "Collecteur"),
        _col("status", "Statut"), _col("items", "Éléments", "int"), _col("message", "Message"),
    ]
    errors = sum(1 for r in rows if r["status"] == "ERROR")
    kpis = [
        {"label": "Exécutions", "value": len(rows), "type": "int"},
        {"label": "Échecs", "value": errors, "type": "int"},
        {"label": "Taux de succès", "value": (1 - errors / len(rows)) if rows else None, "type": "pct"},
    ]
    return _result("collections", "Journal des collectes", p.describe(), columns, rows, kpis=kpis)


def _insight(name: str) -> Callable[[Database, ReportParams], dict]:
    def run(db: Database, p: ReportParams) -> dict:
        from heptagoneslm.web import insights

        return getattr(insights, name)(db, p)

    return run


REPORTS: dict[str, tuple[str, Callable[[Database, ReportParams], dict]]] = {
    "dashboard": ("Tableau de bord", _insight("report_dashboard")),
    "live": ("En direct", _insight("report_live")),
    "summary": ("Synthèse", report_summary),
    "timeseries": ("Usage dans le temps", report_timeseries),
    "breakdown": ("Répartition", report_breakdown),
    "matrix": ("Croisement", report_matrix),
    "heatmap": ("Profil horaire", report_heatmap),
    "contention": ("Contention", report_contention),
    "denials": ("Refus", report_denials),
    "efficiency": ("Efficacité", _insight("report_efficiency")),
    "optimization": ("Optimisation & coûts", _insight("report_optimization")),
    "unused": ("Licences inutilisées", _insight("report_unused")),
    "reservations": ("Réservations", _insight("report_reservations")),
    "sessions": ("Sessions", report_sessions),
    "collections": ("Collectes", report_collections),
}


def run_report(db: Database, report_type: str, p: ReportParams) -> dict:
    if report_type not in REPORTS:
        raise ParamError(f"Rapport inconnu : {report_type}")
    return REPORTS[report_type][1](db, p)


# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------


def _filename(report_type: str, p: ReportParams, ext: str) -> str:
    return f"HeptagonesLM_{report_type}_{p.start:%Y%m%d}_{(p.end - timedelta(seconds=1)):%Y%m%d}.{ext}"


def _csv_value(value: Any, type_: str) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return csv_text(value)  # pas de formule exécutable à l'ouverture dans Excel
    if isinstance(value, float):
        if type_ == "pct":
            value = value * 100
        return f"{value:.4f}".rstrip("0").rstrip(".").replace(".", ",")
    return str(value)


def export_csv(result: dict) -> bytes:
    """CSV pour Excel en français : séparateur « ; », décimale « , », UTF-8 avec BOM.

    Les pourcentages sont exprimés en % (0-100).
    """
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    cols = result["columns"]
    writer.writerow([csv_text(c["label"] + (" (%)" if c["type"] == "pct" else "")) for c in cols])
    for row in result["rows"]:
        writer.writerow([_csv_value(row.get(c["key"]), c["type"]) for c in cols])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def export_json(result: dict) -> bytes:
    return json.dumps(result, ensure_ascii=False, indent=2, default=str).encode("utf-8")


def export_xlsx(result: dict) -> bytes:
    """Classeur d'une feuille reprenant le rapport affiché (+ graphique Excel natif)."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    from heptagoneslm.reporting.excel import _autosize, _show_axes, _write_table

    formats = {"pct": "0.0%", "float": "0.00", "hours": "0.0", "int": "0", "datetime": "dd/mm/yyyy hh:mm",
               "date": "dd/mm/yyyy", "money": '#,##0 "' + result.get("currency", "€").replace('"', "") + '"'}
    wb = Workbook()
    ws = wb.active
    ws.title = result["title"][:31].replace("/", "-").replace(":", "")
    ws["A1"] = f"HeptagonesLM — {result['title']}"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = result["description"]
    ws["A2"].font = Font(italic=True, color="666666")

    cols = result["columns"]

    def cell_value(value: Any, type_: str) -> Any:
        if value is None:
            return None
        if type_ == "datetime":
            return datetime.fromisoformat(value)
        if type_ == "date":
            return date.fromisoformat(value)
        return value

    rows = [[cell_value(r.get(c["key"]), c["type"]) for c in cols] for r in result["rows"]]
    fmt = {i + 1: formats[c["type"]] for i, c in enumerate(cols) if c["type"] in formats}
    top, last = _write_table(ws, 4, [c["label"] for c in cols], rows, "tRapport", fmt)
    ws.freeze_panes = "B5"
    for note_row, note in enumerate(result.get("notes", []), start=last + 2):
        ws.cell(row=note_row, column=1, value=note).font = Font(italic=True, color="666666")
    _autosize(ws)

    chart = result.get("chart")
    if chart and chart["kind"] == "heatmap" and rows:
        ws.conditional_formatting.add(
            f"B{top + 1}:H{last}",
            ColorScaleRule(start_type="num", start_value=0, start_color="F8F8F8",
                           end_type="num", end_value=1, end_color="2A78D6"),
        )
    elif chart and chart["kind"] == "line" and chart["series"]:
        cs = wb.create_sheet("Graphique")
        cs.append([result["columns"][0]["label"]] + [s["name"] for s in chart["series"]])
        excel_row(cs)
        for i, x in enumerate(chart["x"]):
            cs.append([x] + [s["values"][i] for s in chart["series"]])
        lc = LineChart()
        lc.title = chart["title"]
        lc.height, lc.width = 10, 26
        if chart["format"] == "pct":
            lc.y_axis.number_format = "0%"
        lc.add_data(Reference(cs, min_col=2, max_col=1 + len(chart["series"]), min_row=1, max_row=1 + len(chart["x"])),
                    titles_from_data=True)
        lc.set_categories(Reference(cs, min_col=1, min_row=2, max_row=1 + len(chart["x"])))
        _show_axes(lc)
        cs.add_chart(lc, f"{get_column_letter(len(chart['series']) + 3)}2")
    elif chart and chart["kind"] == "bar" and chart["items"]:
        cs = wb.create_sheet("Graphique")
        cs.append(["Libellé", chart["title"]])
        for item in chart["items"]:
            cs.append([item["label"], item["value"]])
            excel_row(cs)
        bc = BarChart()
        bc.type = "bar"
        bc.title = chart["title"]
        bc.legend = None
        bc.x_axis.scaling.orientation = "maxMin"
        if chart["format"] == "pct":
            bc.y_axis.number_format = "0%"
        n = len(chart["items"])
        bc.add_data(Reference(cs, min_col=2, min_row=1, max_row=1 + n), titles_from_data=True)
        bc.set_categories(Reference(cs, min_col=1, min_row=2, max_row=1 + n))
        bc.height, bc.width = max(7.5, 0.6 * n + 3), 20
        _show_axes(bc)
        cs.add_chart(bc, "D2")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_full_workbook(db: Database, p: ReportParams, top_n: int = 10) -> bytes:
    """Classeur complet (8 onglets) avec les filtres serveurs/produits appliqués."""
    from heptagoneslm.reporting.excel import build_workbook

    buf = io.BytesIO()
    build_workbook(_data(db, p), top_n).save(buf)
    return buf.getvalue()


EXPORT_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
    "json": "application/json; charset=utf-8",
}


def export(db: Database, report_type: str, p: ReportParams, fmt: str) -> tuple[bytes, str, str]:
    """Retourne (contenu, nom de fichier, type MIME)."""
    if report_type == "full":
        if p.scoped:
            raise PermissionError("Le classeur complet est réservé aux profils à vue globale : "
                                  "exporter les vues une par une.")
        return export_full_workbook(db, p), _filename("complet", p, "xlsx"), EXPORT_TYPES["xlsx"]
    if fmt not in EXPORT_TYPES:
        raise ParamError(f"Format d'export inconnu : {fmt}")
    if report_type == "dashboard":  # pas de tableau propre : on exporte la synthèse
        report_type = "summary"
    result = run_report(db, report_type, p)
    content = {"xlsx": export_xlsx, "csv": export_csv, "json": export_json}[fmt](result)
    return content, _filename(report_type, p, fmt), EXPORT_TYPES[fmt]
