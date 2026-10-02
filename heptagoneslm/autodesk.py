"""Licences Autodesk « utilisateur nommé » : import des rapports d'usage du portail Autodesk Account.

Depuis 2021, la plupart des abonnements Autodesk sont nominatifs : pas de serveur de licences,
l'usage n'est connu que d'Autodesk. L'administrateur l'exporte depuis le portail
(manage.autodesk.com › Reporting › Usage report › Export, au format CSV ou Excel) ; HeptagonesLM
importe ces exports et repère les licences affectées mais inutilisées, à réaffecter ou à ne pas
renouveler.

Sources acceptées :

* commande ``heptagoneslm import-autodesk export.csv`` (ou ``.xlsx``) ;
* dossier surveillé (``autodesk.import_dir``) : chaque export déposé est importé à la collecte
  suivante, une seule fois (empreinte du fichier).

Les noms de colonnes sont reconnus en anglais (``product_name``, ``days_used``… et la nouvelle
présentation 2026 « Product name », « Days used »…) et en français. Une licence est inactive
lorsque l'utilisateur ne l'a pas lancée depuis ``inactive_days`` jours (30 par défaut).
"""

from __future__ import annotations

import csv
import fnmatch
import hashlib
import io
import logging
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from heptagoneslm.storage.db import Database

log = logging.getLogger(__name__)

VENDOR = "autodesk"
EXTENSIONS = (".csv", ".xlsx", ".txt")

# Colonnes reconnues (noms normalisés : minuscules, sans accents, « _ » pour les séparateurs)
ALIASES: dict[str, tuple[str, ...]] = {
    "email": ("email", "user_email", "e_mail", "email_address", "adresse_e_mail", "adresse_email", "courriel",
              "user_principal_name", "upn"),
    "product": ("product_name", "product", "offering_name", "offering", "produit", "nom_du_produit"),
    "first_name": ("first_name", "prenom", "given_name"),
    "last_name": ("last_name", "nom", "surname", "family_name"),
    "name": ("user_name", "name", "full_name", "display_name", "nom_complet", "utilisateur"),
    "assigned_date": ("assigned_date", "assignment_date", "date_assigned", "seat_assigned_date",
                      "date_d_affectation", "date_affectation"),
    "unassigned_date": ("unassigned_date", "date_unassigned", "date_de_desaffectation"),
    "days_used": ("days_used", "jours_d_utilisation", "jours_utilises", "nombre_de_jours_d_utilisation"),
    "days_inactive": ("days_inactive", "jours_d_inactivite", "jours_inactifs"),
    "last_used": ("last_used", "last_used_date", "last_usage_date", "last_accessed", "last_access_date",
                  "last_activity", "last_activity_date", "derniere_utilisation", "date_de_derniere_utilisation"),
    "seat_assignment": ("seat_assignment", "assignment_status", "statut_d_affectation"),
    "user_activity": ("user_activity", "activity", "activite"),
    "version": ("version", "product_version"),
    "report_date": ("report_date", "export_date", "date_du_rapport", "date"),
}
REQUIRED = ("email", "product")


class ImportError_(ValueError):
    """Export illisible (colonnes manquantes, format inconnu)."""


@dataclass
class SeatRow:
    product: str
    email: str
    name: str | None = None
    assigned_date: date | None = None
    unassigned_date: date | None = None
    days_used: int | None = None
    days_inactive: int | None = None
    last_used: date | None = None
    version: str | None = None
    activity: str | None = None


@dataclass
class ParsedReport:
    rows: list[SeatRow]
    report_date: date
    columns: dict[str, str] = field(default_factory=dict)   # champ -> colonne d'origine
    skipped: int = 0


# ------------------------------------------------------------------------------------- lecture


def _norm(header: str) -> str:
    text = unicodedata.normalize("NFKD", str(header or "")).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", text.strip().lower()).strip("_")


def map_columns(headers: list[str]) -> dict[str, int]:
    normalized = [_norm(h) for h in headers]
    out: dict[str, int] = {}
    for key, names in ALIASES.items():
        for name in names:
            if name in normalized:
                out[key] = normalized.index(name)
                break
    return out


def parse_date(value, dayfirst: bool | None = None) -> date | None:
    """Date d'un export ; ``dayfirst`` tranche les dates ambiguës (jj/mm ou mm/jj, 10/02/2026)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text or text.lower() in ("-", "n/a", "na", "none", "never", "jamais"):
        return None
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r"^(\d{1,2})[/.](\d{1,2})[/.](\d{2,4})", text)
    if m:
        a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        y += 2000 if y < 100 else 0
        if a > 12 or "." in text:
            day, month = a, b
        elif b > 12:
            day, month = b, a
        else:
            day, month = (a, b) if dayfirst is not False else (b, a)
        try:
            return date(y, month, day)
        except ValueError:
            return None
    for fmt in ("%d %b %Y", "%b %d, %Y", "%d %B %Y", "%B %d, %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _int(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(str(value).replace(",", ".").strip()))
    except ValueError:
        return None


def _read_table(path: Path, data: bytes) -> list[list]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    text = None
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ImportError_("Encodage du fichier non reconnu")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


DATE_FIELDS = ("assigned_date", "unassigned_date", "last_used", "report_date")
FRENCH_HEADERS = {"produit", "nom_du_produit", "adresse_e_mail", "adresse_email", "courriel", "prenom", "nom",
                  "date_d_affectation", "date_affectation", "derniere_utilisation", "date_de_derniere_utilisation",
                  "jours_d_utilisation", "jours_utilises", "date_du_rapport"}


def _infer_dayfirst(rows: list[list], indexes: list[int], french: bool) -> bool:
    """Ordre jour/mois des dates de l'export : d'après une date non ambiguë, sinon la langue des colonnes.

    Le portail Autodesk en anglais exporte au format américain (mm/jj/aaaa), en français jj/mm/aaaa.
    """
    for row in rows:
        for i in indexes:
            m = re.match(r"^\s*(\d{1,2})/(\d{1,2})/\d{2,4}", str(row[i] if i < len(row) else "") or "")
            if m:
                if int(m.group(1)) > 12:
                    return True
                if int(m.group(2)) > 12:
                    return False
    return french


def parse_report(path: str | Path, data: bytes | None = None, today: date | None = None) -> ParsedReport:
    """Lit un export « Usage report » du portail Autodesk (CSV ou Excel)."""
    path = Path(path)
    data = data if data is not None else path.read_bytes()
    table = [r for r in _read_table(path, data) if any(str(c or "").strip() for c in r)]
    # L'en-tête peut être précédé de lignes de titre : première ligne contenant e-mail et produit
    header_at = next((i for i, r in enumerate(table[:15]) if all(k in map_columns([str(c or "") for c in r])
                                                                for k in REQUIRED)), None)
    if header_at is None:
        raise ImportError_("Colonnes « e-mail » et « produit » introuvables : exporter le rapport « Usage report » "
                           "depuis Autodesk Account (Reporting › Usage report › Export)")
    headers = [str(c or "") for c in table[header_at]]
    cols = map_columns(headers)

    def get(row, key):
        i = cols.get(key)
        return row[i] if i is not None and i < len(row) else None

    dayfirst = _infer_dayfirst(table[header_at + 1:], [cols[k] for k in DATE_FIELDS if k in cols],
                               french=any(_norm(h) in FRENCH_HEADERS for h in headers))
    rows, skipped, report_dates = [], 0, []
    for raw in table[header_at + 1:]:
        email = str(get(raw, "email") or "").strip().lower()
        product = str(get(raw, "product") or "").strip()
        if not email or "@" not in email or not product:
            skipped += 1
            continue
        name = str(get(raw, "name") or "").strip() or " ".join(
            p for p in (str(get(raw, "first_name") or "").strip(), str(get(raw, "last_name") or "").strip()) if p)
        rows.append(SeatRow(
            product=product, email=email, name=name or None,
            assigned_date=parse_date(get(raw, "assigned_date"), dayfirst),
            unassigned_date=parse_date(get(raw, "unassigned_date"), dayfirst),
            days_used=_int(get(raw, "days_used")), days_inactive=_int(get(raw, "days_inactive")),
            last_used=parse_date(get(raw, "last_used"), dayfirst), version=str(get(raw, "version") or "").strip() or None,
            activity=str(get(raw, "user_activity") or "").strip() or None))
        rd = parse_date(get(raw, "report_date"), dayfirst)
        if rd:
            report_dates.append(rd)
    if not rows:
        raise ImportError_("Aucune ligne exploitable dans l'export")
    m = re.search(r"(20\d{2})[-_]?(\d{2})[-_]?(\d{2})", path.name)
    from_name = parse_date(f"{m.group(1)}-{m.group(2)}-{m.group(3)}") if m else None
    report_date = max(report_dates) if report_dates else from_name or today or date.today()
    return ParsedReport(rows, report_date, {k: headers[i] for k, i in cols.items()}, skipped)


# --------------------------------------------------------------------------------------- import


def import_report(db: Database, path: str | Path, data: bytes | None = None, today: date | None = None,
                  force: bool = False) -> dict:
    """Importe un export en base ; un fichier déjà importé (même contenu) est ignoré."""
    path = Path(path)
    data = data if data is not None else path.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    if not force and db.query("SELECT 1 FROM named_import WHERE file_hash = ?", [digest]):
        return {"file": path.name, "status": "déjà importé", "rows": 0}
    parsed = parse_report(path, data, today)
    now = datetime.now().replace(microsecond=0)
    with db.transaction():
        db.execute("DELETE FROM named_seat WHERE vendor = ? AND report_date = ?", [VENDOR, parsed.report_date])
        db.executemany(
            "INSERT INTO named_seat (vendor, product, email, display_name, assigned_date, unassigned_date, days_used, "
            "days_inactive, last_used, version, activity, report_date, imported_at, source_file) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            [(VENDOR, r.product[:255], r.email[:255], (r.name or None) and r.name[:255], r.assigned_date,
              r.unassigned_date, r.days_used, r.days_inactive, r.last_used, r.version and r.version[:50],
              r.activity and r.activity[:50], parsed.report_date, now, path.name[:255]) for r in _dedupe(parsed.rows)])
        db.execute("DELETE FROM named_import WHERE file_hash = ?", [digest])   # réimport forcé
        db.execute("INSERT INTO named_import (vendor, file_name, file_hash, report_date, rows_count, imported_at) "
                   "VALUES (?, ?, ?, ?, ?, ?)", [VENDOR, path.name[:255], digest, parsed.report_date, len(parsed.rows), now])
    log.info("Autodesk : %s importé (%d licences nommées, rapport du %s)", path.name, len(parsed.rows),
             parsed.report_date.strftime("%d/%m/%Y"))
    return {"file": path.name, "status": "importé", "rows": len(parsed.rows), "skipped": parsed.skipped,
            "report_date": parsed.report_date.isoformat(), "columns": parsed.columns}


def _dedupe(rows: list[SeatRow]) -> list[SeatRow]:
    """Une ligne par (produit, utilisateur) : la plus récente utilisation l'emporte."""
    best: dict[tuple, SeatRow] = {}
    for r in rows:
        key = (r.product.lower(), r.email)
        cur = best.get(key)
        if cur is None or (r.last_used or date.min) > (cur.last_used or date.min):
            best[key] = r
    return list(best.values())


def import_folder(db: Database, folder: str | Path) -> list[dict]:
    """Importe les exports nouveaux d'un dossier surveillé (appelé à chaque collecte)."""
    folder = Path(folder)
    if not folder.is_dir():
        return []
    results = []
    for path in sorted(folder.iterdir()):
        if not path.is_file() or path.suffix.lower() not in EXTENSIONS or path.name.startswith(("~$", ".")):
            continue
        try:
            results.append(import_report(db, path))
        except (ImportError_, OSError, ValueError) as exc:
            log.error("Autodesk : %s non importé : %s", path.name, exc)
            results.append({"file": path.name, "status": "erreur", "error": str(exc), "rows": 0})
    return results


# --------------------------------------------------------------------------------------- analyse


def latest_report_date(db: Database) -> date | None:
    value = db.scalar("SELECT MAX(report_date) FROM named_seat WHERE vendor = ?", [VENDOR])
    return parse_date(value) if value else None


def seat_status(row: dict, report_date: date, inactive_days: int) -> tuple[str, int | None]:
    """(statut, jours d'inactivité) : active | inactive | never | unassigned."""
    if row.get("unassigned_date") and parse_date(row["unassigned_date"]) <= report_date:
        return "unassigned", None
    days = row.get("days_inactive")
    last = parse_date(row.get("last_used"))
    if days is None and last:
        days = (report_date - last).days
    if (row.get("days_used") == 0 and not last) or (days is None and row.get("days_used") == 0):
        return "never", days
    if days is None:
        return "active" if (row.get("days_used") or 0) > 0 else "never", None
    return ("inactive" if days >= inactive_days else "active"), days


def annual_cost(product: str, costs: dict[str, float]) -> float | None:
    for pattern, cost in costs.items():
        if fnmatch.fnmatch(product.lower(), str(pattern).lower()):
            return cost
    return None


def seats(db: Database, inactive_days: int = 30, costs: dict[str, float] | None = None,
          products: list[str] | None = None) -> tuple[date | None, list[dict]]:
    """Licences nommées du dernier rapport importé, avec leur statut."""
    report_date = latest_report_date(db)
    if report_date is None:
        return None, []
    rows = db.query_dicts(
        "SELECT product, email, display_name, assigned_date, unassigned_date, days_used, days_inactive, last_used, "
        "version, activity FROM named_seat WHERE vendor = ? AND report_date = ?", [VENDOR, report_date])
    costs = costs or {}
    out = []
    for r in rows:
        if products and not any(p.lower() in r["product"].lower() for p in products):
            continue
        status, days = seat_status(r, report_date, inactive_days)
        if status == "unassigned":
            continue
        unit = annual_cost(r["product"], costs)
        out.append({**r, "status": status, "inactive_for": days, "unit_cost": unit,
                    "recoverable_cost": unit if status in ("inactive", "never") else None})
    order = {"never": 0, "inactive": 1, "active": 2}
    out.sort(key=lambda r: (order[r["status"]], -(r["inactive_for"] or 0), r["product"], r["email"]))
    return report_date, out
