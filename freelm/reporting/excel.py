"""Génération du classeur Excel de reporting (openpyxl).

Feuilles produites :

1. **Synthèse**         : indicateurs par produit/feature + graphique des taux d'utilisation.
2. **Usage quotidien**  : tableau croisé date × produit (pic de jetons) + courbes.
3. **Utilisateurs**     : tableau croisé utilisateur × produit (heures) + top utilisateurs.
4. **Profil horaire**   : carte de chaleur heure × jour de semaine (taux moyen).
5. **Contention**       : jours avec saturation et/ou refus, par produit.
6. **Refus**            : détail des refus (DENIED) issus des logs.
7. **Sessions**         : données brutes au format Tableau Excel, prêtes pour
   un tableau croisé dynamique (Insertion > Tableau croisé dynamique).
8. **À propos**         : période, filtres, définitions des indicateurs.

openpyxl ne sait pas créer de véritables tableaux croisés dynamiques : les
croisements usuels sont précalculés, et les données brutes sont fournies
sous forme de Tableaux Excel pour les analyses ad hoc.
"""

from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from freelm.reporting.dataset import WEEKDAYS_FR, ReportData
from freelm.reporting.sanitize import excel_text

PCT = "0.0%"
DEC1 = "0.0"
DEC2 = "0.00"
DATE = "dd/mm/yyyy"
DATETIME = "dd/mm/yyyy hh:mm"
TABLE_STYLE = "TableStyleMedium2"

TITLE_FONT = Font(size=14, bold=True)
SUBTITLE_FONT = Font(italic=True, color="666666")


def _title(ws: Worksheet, title: str, data: ReportData) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        f"Période du {data.start:%d/%m/%Y %H:%M} au {data.end:%d/%m/%Y %H:%M} (exclu)"
        + (f" — serveurs : {', '.join(data.servers)}" if data.servers else "")
    )
    ws["A2"].font = SUBTITLE_FONT


def _write_table(
    ws: Worksheet,
    top: int,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    name: str,
    formats: dict[int, str] | None = None,
) -> tuple[int, int]:
    """Écrit un Tableau Excel à partir de la ligne ``top``. Retourne (1re, dernière ligne)."""
    formats = formats or {}
    for c, h in enumerate(headers, start=1):
        excel_text(ws.cell(row=top, column=c, value=h))
    for r, row in enumerate(rows, start=top + 1):
        for c, value in enumerate(row, start=1):
            if isinstance(value, str):  # caractères de contrôle éventuels des logs
                value = ILLEGAL_CHARACTERS_RE.sub("", value)
            cell = ws.cell(row=r, column=c, value=value)
            excel_text(cell)  # jamais de formule issue des données
            if c in formats:
                cell.number_format = formats[c]
    last = top + len(rows)
    if rows:
        ref = f"A{top}:{get_column_letter(len(headers))}{last}"
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name=TABLE_STYLE, showRowStripes=True)
        ws.add_table(table)
    else:
        for c in range(1, len(headers) + 1):
            ws.cell(row=top, column=c).font = Font(bold=True)
        ws.cell(row=top + 1, column=1, value="Aucune donnée sur la période").font = SUBTITLE_FONT
    return top, last


def _autosize(ws: Worksheet, min_width: int = 8, max_width: int = 45) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            if cell.value is None:
                continue
            if isinstance(cell.value, datetime):
                length = 16
            elif isinstance(cell.value, float):
                length = 8
            else:
                length = len(str(cell.value))
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, width + 2))


def _show_axes(chart) -> None:
    # openpyxl >= 3.1 masque les axes par défaut dans les versions récentes d'Excel.
    chart.x_axis.delete = False
    chart.y_axis.delete = False


# ---------------------------------------------------------------------------
# Feuilles
# ---------------------------------------------------------------------------


def _sheet_summary(wb: Workbook, data: ReportData, top_n: int) -> None:
    ws = wb.active
    ws.title = "Synthèse"
    _title(ws, "FreeLM — Synthèse de l'utilisation des licences", data)

    summaries = sorted(data.summaries, key=lambda s: (-(s.avg_utilization or 0), -s.usage_hours))
    headers = [
        "Serveur", "Éditeur", "Produit", "Feature", "Jetons émis", "Pic", "Moyenne",
        "Taux moyen", "Taux au pic", "Taux heures ouvrées", "Heures saturées", "Refus",
        "Utilisateurs", "Sessions", "Heures d'utilisation",
    ]
    rows = [
        [
            s.server, s.publisher, s.label, s.feature, s.issued, s.peak_in_use, round(s.avg_in_use, 2),
            s.avg_utilization, s.peak_utilization, s.business_utilization, s.saturated_hours,
            s.denials, s.distinct_users, s.sessions, round(s.usage_hours, 1),
        ]
        for s in summaries
    ]
    top, last = _write_table(
        ws, 4, headers, rows, "tSynthese",
        {7: DEC2, 8: PCT, 9: PCT, 10: PCT, 15: DEC1},
    )
    ws.freeze_panes = "E5"
    if not rows:
        _autosize(ws)
        return

    ws.conditional_formatting.add(
        f"H{top + 1}:H{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="5B9BD5")
    )
    ws.conditional_formatting.add(
        f"I{top + 1}:I{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="ED7D31")
    )

    n = min(top_n, len(rows))
    chart = BarChart()
    chart.type = "bar"
    chart.title = f"Taux d'utilisation — top {n}"
    chart.y_axis.number_format = "0%"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 1
    chart.x_axis.scaling.orientation = "maxMin"  # le plus utilisé en haut
    chart.add_data(Reference(ws, min_col=8, max_col=9, min_row=top, max_row=top + n), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=3, min_row=top + 1, max_row=top + n))
    chart.height = max(7.5, 0.6 * n + 3)
    chart.width = 22
    _show_axes(chart)
    ws.add_chart(chart, f"B{last + 3}")
    _autosize(ws)


def _sheet_daily(wb: Workbook, data: ReportData, top_n: int) -> None:
    ws = wb.create_sheet("Usage quotidien")
    _title(ws, "Pic quotidien de jetons utilisés par produit", data)

    # Colonnes triées par volume d'usage : les N premières alimentent le graphique.
    ordered = sorted(data.summaries, key=lambda s: (-s.usage_hours, -s.peak_in_use))
    labels = [s.label for s in ordered]
    days = []
    d = data.start.date()
    while datetime.combine(d, datetime.min.time()) < data.end:
        days.append(d)
        d += timedelta(days=1)
    rows = [[day] + [data.daily_peak.get(day, {}).get(lab) for lab in labels] for day in days]
    top, last = _write_table(ws, 4, ["Date", *labels], rows, "tQuotidien", {1: DATE})
    ws.freeze_panes = "B5"
    if rows and labels:
        n = min(top_n, len(labels))
        chart = LineChart()
        chart.title = f"Pic quotidien de jetons utilisés — top {n}"
        chart.y_axis.title = "Jetons"
        chart.x_axis.number_format = "dd/mm"
        chart.add_data(Reference(ws, min_col=2, max_col=1 + n, min_row=top, max_row=last), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=top + 1, max_row=last))
        chart.height = 10
        chart.width = 26
        _show_axes(chart)
        ws.add_chart(chart, f"{get_column_letter(len(labels) + 3)}4")
    _autosize(ws)


def _sheet_users(wb: Workbook, data: ReportData, top_n: int) -> None:
    ws = wb.create_sheet("Utilisateurs")
    _title(ws, "Heures d'utilisation par utilisateur et par produit", data)
    products = data.products
    headers = ["Utilisateur", "Nom", "Service", *products, "Total (h)"]
    rows = []
    for user, per in data.user_hours.items():  # déjà trié par total décroissant
        info = data.user_info.get(user, {})
        rows.append(
            [user, info.get("display_name"), info.get("department")]
            + [round(per[p], 2) if p in per else None for p in products]
            + [round(sum(per.values()), 2)]
        )
    total_col = len(headers)
    fmts = {c: DEC1 for c in range(4, total_col + 1)}
    top, last = _write_table(ws, 4, headers, rows, "tUtilisateurs", fmts)
    ws.freeze_panes = "D5"
    if rows:
        for c in range(4, total_col):
            col = get_column_letter(c)
            ws.conditional_formatting.add(
                f"{col}{top + 1}:{col}{last}",
                ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="63BE7B"),
            )
        n = min(top_n, len(rows))
        chart = BarChart()
        chart.type = "bar"
        chart.title = f"Top {n} utilisateurs (heures)"
        chart.x_axis.scaling.orientation = "maxMin"
        chart.legend = None
        chart.add_data(
            Reference(ws, min_col=total_col, min_row=top, max_row=top + n), titles_from_data=True
        )
        chart.set_categories(Reference(ws, min_col=1, min_row=top + 1, max_row=top + n))
        chart.height = max(7.5, 0.6 * n + 3)
        chart.width = 18
        _show_axes(chart)
        ws.add_chart(chart, f"{get_column_letter(total_col + 2)}4")
    _autosize(ws)


def _sheet_hour_profile(wb: Workbook, data: ReportData) -> None:
    ws = wb.create_sheet("Profil horaire")
    _title(ws, "Taux d'utilisation moyen par heure et jour de semaine (tous produits)", data)
    ws.cell(row=4, column=1, value="Heure").font = Font(bold=True)
    for c, day in enumerate(WEEKDAYS_FR, start=2):
        cell = ws.cell(row=4, column=c, value=day)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for hour in range(24):
        r = 5 + hour
        ws.cell(row=r, column=1, value=f"{hour:02d}h")
        for wd in range(7):
            value = data.hour_profile.get((wd, hour))
            cell = ws.cell(row=r, column=2 + wd, value=value)
            cell.number_format = PCT
    ws.conditional_formatting.add(
        "B5:H28",
        ColorScaleRule(
            start_type="num", start_value=0, start_color="F8F8F8",
            mid_type="num", mid_value=0.5, mid_color="FFEB84",
            end_type="num", end_value=1, end_color="F8696B",
        ),
    )
    ws.freeze_panes = "B5"
    for c in range(2, 9):
        ws.column_dimensions[get_column_letter(c)].width = 11


def _sheet_contention(wb: Workbook, data: ReportData) -> None:
    ws = wb.create_sheet("Contention")
    _title(ws, "Jours de contention : saturation des jetons et refus", data)
    headers = ["Date", "Serveur", "Produit", "Feature", "Jetons émis", "Pic", "Heures saturées", "Refus"]
    rows = [
        [c["date"], c["server"], c["product"], c["feature"], c["issued"], c["peak_in_use"],
         c["saturated_hours"], c["denials"]]
        for c in data.contention
    ]
    top, last = _write_table(ws, 4, headers, rows, "tContention", {1: DATE})
    if rows:
        for col in ("G", "H"):
            ws.conditional_formatting.add(
                f"{col}{top + 1}:{col}{last}",
                ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="F8696B"),
            )
    ws.freeze_panes = "A5"
    _autosize(ws)


def _sheet_denials(wb: Workbook, data: ReportData) -> None:
    ws = wb.create_sheet("Refus")
    _title(ws, "Détail des refus de licence (DENIED)", data)
    headers = ["Horodatage", "Serveur", "Produit", "Feature", "Utilisateur", "Poste", "Motif"]
    rows = [
        [d["event_time"], d["server"], d["product"], d["feature"], d["username"], d["hostname"], d["detail"]]
        for d in data.denials
    ]
    _write_table(ws, 4, headers, rows, "tRefus", {1: DATETIME})
    ws.freeze_panes = "A5"
    _autosize(ws)


def _sheet_sessions(wb: Workbook, data: ReportData) -> None:
    ws = wb.create_sheet("Sessions")
    _title(ws, "Sessions de licence (données brutes pour tableaux croisés dynamiques)", data)
    headers = [
        "Serveur", "Éditeur", "Produit", "Feature", "Utilisateur", "Nom", "Service", "Poste",
        "Début", "Fin", "Jetons", "Heures (période)", "Source", "Motif de fin",
    ]
    rows = [
        [s["server"], s["publisher"], s["product"], s["feature"], s["username"], s["display_name"],
         s["department"], s["hostname"], s["checkout_time"], s["checkin_time"], s["licenses"],
         s["hours_in_period"], s["source"], s["close_reason"]]
        for s in data.sessions
    ]
    _write_table(ws, 4, headers, rows, "tSessions", {9: DATETIME, 10: DATETIME, 12: DEC2})
    ws.freeze_panes = "A5"
    _autosize(ws)


def _sheet_about(wb: Workbook, data: ReportData) -> None:
    ws = wb.create_sheet("À propos")
    ws["A1"] = "FreeLM — rapport d'utilisation des licences"
    ws["A1"].font = TITLE_FONT
    lines = [
        ("Généré le", data.generated_at.strftime("%d/%m/%Y %H:%M")),
        ("Début de période", data.start.strftime("%d/%m/%Y %H:%M")),
        ("Fin de période (exclue)", data.end.strftime("%d/%m/%Y %H:%M")),
        ("Serveurs", ", ".join(data.servers) or "tous"),
        ("", ""),
        ("Jetons émis", "Nombre maximal de jetons déclarés par le serveur sur la période."),
        ("Pic", "Nombre maximal de jetons utilisés simultanément (sondages lmstat)."),
        ("Moyenne", "Nombre moyen de jetons utilisés sur l'ensemble des sondages."),
        ("Taux moyen / au pic", "Moyenne / pic rapportés aux jetons émis."),
        ("Taux heures ouvrées", "Taux moyen restreint aux jours ouvrés et aux heures ouvrées configurées."),
        ("Heures saturées", "Heures durant lesquelles tous les jetons ont été utilisés au moins une fois."),
        ("Refus", "Demandes refusées (DENIED) relevées dans les logs du serveur."),
        ("Heures d'utilisation", "Somme des durées de session × jetons, tronquées à la période."),
        ("Précision", "Sessions issues de lmstat : ± l'intervalle de sondage ; issues des logs : à la seconde."),
    ]
    for r, (k, v) in enumerate(lines, start=3):
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100


def build_workbook(data: ReportData, top_n: int = 10) -> Workbook:
    """Construit le classeur complet (8 feuilles) en mémoire."""
    wb = Workbook()
    _sheet_summary(wb, data, top_n)
    _sheet_daily(wb, data, top_n)
    _sheet_users(wb, data, top_n)
    _sheet_hour_profile(wb, data)
    _sheet_contention(wb, data)
    _sheet_denials(wb, data)
    _sheet_sessions(wb, data)
    _sheet_about(wb, data)
    return wb


def write_excel_report(data: ReportData, output: str | Path, top_n: int = 10) -> Path:
    """Écrit le classeur Excel et retourne son chemin."""
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(data, top_n).save(output)
    return output
