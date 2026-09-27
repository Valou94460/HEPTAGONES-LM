import importlib.util
import os
import tempfile
import unittest
from datetime import date, datetime, timedelta

from freelm.reporting.dataset import build_report_data, previous_month
from freelm.service import CollectionService
from tests.helpers import make_config, make_db
from tests.test_collection import co, status_at

HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None
DAY = datetime(2024, 9, 23)  # lundi


class ReportDatasetTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg = make_config()
        cls.db = make_db()
        svc = CollectionService(cls.cfg, cls.db)
        sid = svc.sync_servers()["AUTODESK"]
        # 9h-11h : alice seule ; 11h-12h : alice + bob (2 jetons) sur 3 émis -> saturation
        a = co("alice", "1", DAY.replace(hour=9))
        b = co("bob", "2", DAY.replace(hour=11), count=2)
        t = DAY.replace(hour=9)
        while t < DAY.replace(hour=13):
            cur = [a] + ([b] if DAY.replace(hour=11) <= t < DAY.replace(hour=12) else [])
            with cls.db.transaction():
                svc.apply_status(sid, status_at(t, cur, issued=3), "lmstat", True)
            t += timedelta(minutes=15)
        with cls.db.transaction():
            fid = cls.db.scalar("SELECT id FROM feature")
            uid = svc.repo.get_or_create_user("carol")
            hid = svc.repo.get_or_create_host("PC-carol")
            svc.repo.add_event(sid, DAY.replace(hour=11, minute=30), "DENIED", fid, uid, hid, 1, "Licensed number of users already reached")
        cls.data = build_report_data(cls.db, DAY, DAY + timedelta(days=1), now=DAY + timedelta(days=2))

    def test_summary(self):
        (s,) = self.data.summaries
        self.assertEqual(s.label, "AutoCAD 2024")
        self.assertEqual(s.issued, 3)
        self.assertEqual(s.peak_in_use, 3)
        self.assertEqual(s.saturated_hours, 1)
        self.assertEqual(s.denials, 1)
        self.assertEqual(s.distinct_users, 2)
        # 16 sondages : 12 à 1 jeton, 4 à 3 jetons -> moyenne 1.5
        self.assertAlmostEqual(s.avg_in_use, 1.5)
        self.assertAlmostEqual(s.avg_utilization, 0.5)
        self.assertAlmostEqual(s.peak_utilization, 1.0)
        # alice : session toujours ouverte, vue jusqu'à 12h45 -> 3h45 ; bob : 2 x 52,5 min
        self.assertAlmostEqual(s.usage_hours, 3.75 + 1.75)

    def test_users_and_profile(self):
        self.assertEqual(list(self.data.user_hours)[0], "alice")
        self.assertAlmostEqual(self.data.user_hours["bob"]["AutoCAD 2024"], 2 * (52.5 / 60))  # vu de 11h00 à 11h45, fin estimée 11h52m30
        self.assertAlmostEqual(self.data.hour_profile[(0, 11)], 3 / 3)
        self.assertAlmostEqual(self.data.hour_profile[(0, 9)], 1 / 3)
        self.assertEqual(self.data.daily_peak[DAY.date()]["AutoCAD 2024"], 3)

    def test_contention(self):
        (row,) = self.data.contention
        self.assertEqual((row["saturated_hours"], row["denials"]), (1, 1))

    def test_server_filter(self):
        data = build_report_data(self.db, DAY, DAY + timedelta(days=1), servers=["AUTRE"])
        self.assertEqual(data.summaries, [])

    def test_previous_month(self):
        self.assertEqual(previous_month(date(2024, 1, 15)), (datetime(2023, 12, 1), datetime(2024, 1, 1)))

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl non installé")
    def test_excel_export(self):
        from openpyxl import load_workbook

        from freelm.reporting.excel import write_excel_report

        with tempfile.TemporaryDirectory() as tmp:
            path = write_excel_report(self.data, os.path.join(tmp, "r.xlsx"))
            wb = load_workbook(path)
            self.assertEqual(
                wb.sheetnames,
                ["Synthèse", "Usage quotidien", "Utilisateurs", "Profil horaire", "Contention", "Refus", "Sessions", "À propos"],
            )
            self.assertEqual(wb["Synthèse"]["C5"].value, "AutoCAD 2024")

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl non installé")
    def test_excel_export_empty_period(self):
        from freelm.reporting.excel import write_excel_report

        empty = build_report_data(self.db, datetime(2020, 1, 1), datetime(2020, 2, 1))
        with tempfile.TemporaryDirectory() as tmp:
            self.assertTrue(write_excel_report(empty, os.path.join(tmp, "e.xlsx")).exists())


if __name__ == "__main__":
    unittest.main()
