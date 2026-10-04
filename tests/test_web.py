import importlib.util
import json
import os
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from datetime import date, datetime, timedelta

from heptagoneslm.config import load_config
from heptagoneslm.simulation.demo import run_demo
from heptagoneslm.storage import Database, Repository
from heptagoneslm.users import import_users, read_users_csv
from heptagoneslm.web import api
from heptagoneslm.web.server import make_server

HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None
MONDAY = datetime(2024, 9, 23)


class ParamsTest(unittest.TestCase):
    def test_defaults_last_30_days(self):
        p = api.parse_params({}, today=date(2024, 9, 30))
        self.assertEqual(p.start, datetime(2024, 9, 1))
        self.assertEqual(p.end, datetime(2024, 10, 1))  # fin incluse -> lendemain exclu

    def test_lists_and_validation(self):
        p = api.parse_params({"server": ["A", "B"], "user": ["J.Doe"], "start": ["2024-01-01"], "end": ["2024-01-31"]})
        self.assertEqual(p.servers, ["A", "B"])
        self.assertEqual(p.users, ["j.doe"])
        for bad in ({"start": ["2024-02-01"], "end": ["2024-01-01"]}, {"start": ["x"]},
                    {"granularity": ["year"]}, {"dimension": ["planet"]}):
            with self.assertRaises(api.ParamError):
                api.parse_params(bad)


class WebServerTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        result = run_demo(cls.tmp.name, days=7, end=MONDAY + timedelta(days=7), make_report=False)
        cls.cfg = load_config(result.config_path)
        cls.httpd = make_server(cls.cfg, port=0)
        cls.base = f"http://127.0.0.1:{cls.httpd.server_address[1]}"
        threading.Thread(target=cls.httpd.serve_forever, daemon=True).start()
        cls.period = "start=2024-09-16&end=2024-09-29"

    @classmethod
    def tearDownClass(cls):
        cls.httpd.shutdown()
        cls.httpd.server_close()
        cls.tmp.cleanup()

    def get(self, path):
        try:
            with urllib.request.urlopen(self.base + path) as r:
                return r.status, r.headers, r.read()
        except urllib.error.HTTPError as e:
            return e.code, e.headers, e.read()

    def get_json(self, path):
        status, _, body = self.get(path)
        return status, json.loads(body)

    def test_static_files(self):
        for path, ctype in (("/", "text/html"), ("/app.js", "text/javascript"), ("/style.css", "text/css"),
                            ("/theme.js", "text/javascript")):
            status, headers, _ = self.get(path)
            self.assertEqual(status, 200)
            self.assertTrue(headers["Content-Type"].startswith(ctype), headers["Content-Type"])
        self.assertEqual(self.get("/../heptagoneslm.demo.yaml")[0], 404)

    def test_meta(self):
        status, meta = self.get_json("/api/meta")
        self.assertEqual(status, 200)
        self.assertEqual(meta["servers"], ["ANSYS", "AUTODESK", "Autodesk Cloud", "MATLAB"])   # licences nommées
        self.assertIn("Simulation & calcul", meta["departments"])
        self.assertEqual(len(meta["collections"]), 3)  # un collecteur « simulation » par serveur
        self.assertIn("summary", meta["reports"])

    def test_all_reports(self):
        for rtype in api.REPORTS:
            status, r = self.get_json(f"/api/report?type={rtype}&{self.period}&dimension=department")
            self.assertEqual(status, 200, (rtype, r))
            if rtype == "dashboard":  # vue composite : indicateurs, tendances, blocs
                self.assertEqual(len(r["kpis"]), 6)
                self.assertTrue(r["products"] and r["trend"]["x"] and r["servers"])
                continue
            self.assertTrue(r["rows"], rtype)
            for col in r["columns"]:  # chaque colonne annoncée est présente dans les lignes
                self.assertIn(col["key"], r["rows"][0], (rtype, col["key"]))

    def test_insight_reports(self):
        _, eff = self.get_json(f"/api/report?type=efficiency&{self.period}")
        for row in eff["rows"]:
            self.assertLessEqual(row["recommended"], row["issued"], row)
            if row["server"] != "Autodesk Cloud":       # licences nommées : recommandé = licences actives
                self.assertLessEqual(row["recommended"], row["peak"], row)
        self.assertEqual(eff["chart"]["xType"], "share")
        _, one = self.get_json(f"/api/report?type=efficiency&{self.period}&product=AutoCAD%202024")
        self.assertEqual(one["chart"]["xType"], "tokens")
        values = one["chart"]["series"][0]["values"]
        self.assertEqual(values, sorted(values, reverse=True))  # « au moins N » : courbe décroissante

        _, opt = self.get_json(f"/api/report?type=optimization&{self.period}")
        self.assertEqual(opt["currency"], "€")
        costed = [r for r in opt["rows"] if r["unit_cost"] is not None]
        self.assertTrue(costed)  # la démo renseigne annual_cost
        for r in costed:
            self.assertEqual(r["annual_cost"], r["unit_cost"] * r["issued"])
            self.assertEqual(r["savings"], r["surplus"] * r["unit_cost"])
        for r in opt["rows"]:
            if r["denials"]:
                self.assertEqual(r["surplus"], 0, r)  # jamais de baisse conseillée en cas de refus

        _, live = self.get_json("/api/report?type=live")
        self.assertTrue(live["live"]["updated_at"] and live["live"]["products"])
        for p in live["live"]["products"]:
            self.assertEqual(p["free"], p["issued"] - p["in_use"])

        _, dash = self.get_json(f"/api/report?type=dashboard&{self.period}")
        keys = [k["key"] for k in dash["kpis"]]
        self.assertEqual(keys[:5], ["utilization", "users", "hours", "saturated", "denials"])
        self.assertEqual(len(dash["heatmap"]["values"]), 24)
        self.assertEqual(len(dash["trend"]["x"]), len(dash["trend"]["business_utilization"]))
        for a in dash["alerts"]:
            self.assertIn(a["severity"], ("critical", "serious", "warning", "info"))

    def test_filters(self):
        _, all_ = self.get_json(f"/api/report?type=sessions&{self.period}")
        _, one = self.get_json(f"/api/report?type=sessions&{self.period}&server=ANSYS")
        self.assertTrue(0 < len(one["rows"]) < len(all_["rows"]))
        self.assertEqual({r["server"] for r in one["rows"]}, {"ANSYS"})
        user = one["rows"][0]["username"]
        _, mine = self.get_json(f"/api/report?type=sessions&{self.period}&user={user}")
        self.assertEqual({r["username"] for r in mine["rows"]}, {user})
        _, prod = self.get_json(f"/api/report?type=timeseries&{self.period}&product=AutoCAD%202024&granularity=hour")
        self.assertEqual({r["product"] for r in prod["rows"]}, {"AutoCAD 2024"})
        self.assertEqual(len(prod["chart"]["series"]), 1)

    def test_timeseries_granularities(self):
        counts = {}
        for g in api.GRANULARITIES:
            _, r = self.get_json(f"/api/report?type=timeseries&{self.period}&granularity={g}&metric=utilization")
            counts[g] = len(r["chart"]["x"])
        self.assertGreater(counts["hour"], counts["day"])
        self.assertGreaterEqual(counts["day"], counts["week"])
        self.assertGreaterEqual(counts["week"], counts["month"])

    def test_csv_export(self):
        status, headers, body = self.get(f"/api/export?type=breakdown&format=csv&{self.period}&dimension=department")
        self.assertEqual(status, 200)
        self.assertIn("attachment", headers["Content-Disposition"])
        text = body.decode("utf-8-sig")
        self.assertTrue(text.startswith("Service;Heures;Part (%)"))
        self.assertIn(",", text.splitlines()[1])  # décimale française

    def test_json_export(self):
        status, _, body = self.get(f"/api/export?type=summary&format=json&{self.period}")
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(body)["type"], "summary")

    def test_errors(self):
        self.assertEqual(self.get("/api/report?type=nope")[0], 400)
        self.assertEqual(self.get("/api/report?start=2024-13-01")[0], 400)
        self.assertEqual(self.get("/api/export?type=summary&format=pdf")[0], 400)
        self.assertEqual(self.get("/api/unknown")[0], 404)

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl non installé")
    def test_xlsx_exports(self):
        from io import BytesIO

        from openpyxl import load_workbook

        for rtype in list(api.REPORTS) + ["full"]:
            status, headers, body = self.get(f"/api/export?type={rtype}&format=xlsx&{self.period}")
            self.assertEqual(status, 200, rtype)
            wb = load_workbook(BytesIO(body))
            self.assertGreater(wb.worksheets[0].max_row, 4, rtype)


class UsersImportTest(unittest.TestCase):
    def test_ad_export_columns(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "u.csv")
            with open(path, "w", encoding="utf-8-sig") as fh:
                fh.write('"SamAccountName","DisplayName","Department"\n"JDoe","John Doe","CAO"\n"asmith","Anna Smith",""\n')
            users = read_users_csv(path)
            self.assertEqual(users[0], {"username": "JDoe", "display_name": "John Doe", "department": "CAO"})
            db = Database("sqlite:///:memory:")
            from heptagoneslm.storage import create_schema

            create_schema(db)
            repo = Repository(db)
            repo.get_or_create_user("jdoe")
            self.assertEqual(import_users(repo, users), (1, 1))
            self.assertEqual(db.query("SELECT username, display_name, department FROM app_user ORDER BY username"),
                             [("asmith", "Anna Smith", None), ("jdoe", "John Doe", "CAO")])


if __name__ == "__main__":
    unittest.main()


class AlertsTest(unittest.TestCase):
    def test_compute_alerts(self):
        from types import SimpleNamespace

        from heptagoneslm.web import insights
        p = api.parse_params({"start": ["2024-09-16"], "end": ["2024-09-22"]})
        data = SimpleNamespace(summaries=[
            SimpleNamespace(label="A", denials=3, saturated_hours=0, issued=4, peak_utilization=1.0, peak_in_use=4),
            SimpleNamespace(label="B", denials=0, saturated_hours=0, issued=10, peak_utilization=0.3, peak_in_use=3),
        ])
        now = datetime(2024, 9, 22, 12, 0)
        collections = [{"server": "S", "collector": "lmstat", "status": "ERROR", "message": "timeout",
                        "last_run": "2024-09-22 11:55"}]
        expiring = [{"product": "X", "feature": "fx", "server": "S", "expires_on": "2024-09-25", "days_left": 3},
                    {"product": "Y", "feature": "fy", "server": "S", "expires_on": "2024-10-10", "days_left": 18},
                    {"product": "Z", "feature": "fz", "server": "S", "expires_on": "2024-10-12", "days_left": 20}]
        alerts = insights.compute_alerts(None, p, data, collections, expiring, now=now)
        sev = [a["severity"] for a in alerts]
        self.assertEqual(sev, sorted(sev, key=["critical", "serious", "warning", "info"].index))
        titles = " | ".join(a["title"] for a in alerts)
        self.assertIn("Collecte en échec", titles)
        self.assertIn("Expiration dans 3 j", titles)
        self.assertIn("2 licence(s) expirent", titles)  # avertissements regroupés
        self.assertIn("3 refus", titles)
        self.assertIn("sous-utilisé", titles)
        self.assertNotIn("Collecte arrêtée", titles)
        stale = insights.compute_alerts(None, p, data, [dict(collections[0], status="OK")], [],
                                        now=now + timedelta(hours=3))
        self.assertIn("Collecte arrêtée", stale[0]["title"])
